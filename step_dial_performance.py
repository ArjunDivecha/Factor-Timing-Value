"""
=============================================================================
SCRIPT NAME: step_dial_performance.py  (helper module used by Step Ten Exposure Dial.py)
=============================================================================

DESCRIPTION:
    Shows what the exposure dial does to the strategy: the UNLEVERED strategy (always 100%
    invested, i.e. what Step Nine reports) versus the DIAL strategy (exposure 0%..200%).

    How a month is computed (MONTHLY accounting, gross of trading costs):
      exposure for month m = the dial reading on the last trading day of month m-1
      r        = the unlevered strategy's return in month m (Step Nine 'Portfolio' column)
      rf       = T-bill return for that month (yfinance ^IRX, 13-week bill yield)
      exposure L <= 1:  dial return = L * r + (1 - L) * rf          (the rest sits in cash)
      exposure L >  1:  dial return = L * r - (L - 1) * (rf + 0.42%)  (borrowing on margin)
    The 0.42% margin spread over the T-bill is a CONSTANT taken from the research
    (4.5% margin rate vs 4.08% bill, Sept 2026) - it flatters earlier years when real
    spreads were wider. Exposure is held fixed within the month (the research used a daily
    model where leverage drifts with prices; this is the simpler monthly version).
    NO transaction costs: research is judged gross.

INPUT FILES:
    <repo>/T2_Final_Portfolio_Returns.xlsx   sheet 'Monthly Returns', column 'Portfolio'
        (index = month start; the row dated YYYY-MM-01 is the return during that month)
    Yahoo Finance '^IRX' daily 13-week T-bill yield, percent (downloaded at run time)
    the daily exposure history passed in by Step Ten (a pandas DataFrame, column 'exposure')

OUTPUT FILES:
    <repo>/T2_exposure_dial_performance.pdf   (written by plot_performance)
        page: growth of $1 (log) unlevered vs dial, and drawdown of each
    The statistics table is returned to Step Ten, which prints it and stores it in the
    'Dial_Performance' sheet of the FINAL workbook.

VERSION: 1.0 (2026-10-01)
DEPENDENCIES: pandas, numpy, matplotlib, yfinance, openpyxl
=============================================================================
"""
import numpy as np
import pandas as pd

MARGIN_SPREAD = 0.0042          # annual, over the T-bill (research constant)
PERIODS = [("Full sample", None), ("Since 2012 (all 34 ETFs live)", "2012-01-01"),
           ("Last 10 years", -120), ("Last 5 years", -60)]


def load_tbill_monthly():
    """Monthly T-bill return from yfinance ^IRX (annual yield in percent)."""
    import yfinance as yf
    d = yf.download("^IRX", period="max", interval="1d", progress=False, auto_adjust=False)
    if d is None or d.empty:
        raise RuntimeError("yfinance returned no ^IRX (T-bill) data")
    y = d["Close"]
    y = (y.iloc[:, 0] if isinstance(y, pd.DataFrame) else y).dropna()
    y = y[y > -5]
    avg = (y / 100.0).groupby(y.index.to_period("M")).mean()
    rf = (1 + avg) ** (1 / 12) - 1
    rf.index = rf.index.to_timestamp()
    return rf


def build_monthly(exposure_daily, portfolio_path):
    """Month-by-month unlevered vs dial returns. Returns a DataFrame indexed by month start."""
    r = pd.read_excel(portfolio_path, sheet_name="Monthly Returns", index_col=0)["Portfolio"].dropna()
    r.index = pd.to_datetime(r.index).to_period("M").to_timestamp()
    # exposure held in month m = reading at the last trading day of month m-1
    me = exposure_daily.groupby(exposure_daily.index.to_period("M")).last()
    held = me.copy()
    held.index = (held.index + 1).to_timestamp()
    rf = load_tbill_monthly().reindex(r.index).ffill()
    m = pd.DataFrame({"unlevered": r, "exposure": held.reindex(r.index), "rf": rf}).dropna()
    L = m["exposure"]
    spread_m = (1 + MARGIN_SPREAD) ** (1 / 12) - 1
    m["dial"] = np.where(L <= 1.0, L * m["unlevered"] + (1 - L) * m["rf"],
                         L * m["unlevered"] - (L - 1.0) * (m["rf"] + spread_m))
    return m


def _stats(ret, rf):
    cum = (1 + ret).cumprod()
    yrs = len(ret) / 12.0
    ex = ret - rf
    return {"Months": len(ret),
            "Ann. return (CAGR)": cum.iloc[-1] ** (1 / yrs) - 1,
            "Ann. volatility": ret.std() * np.sqrt(12),
            "Sharpe (excess over T-bill)": ex.mean() / ex.std() * np.sqrt(12),
            "Max drawdown": (cum / cum.cummax() - 1).min(),
            "Worst month": ret.min(), "Best month": ret.max(),
            "Growth of $1": cum.iloc[-1]}


def performance_table(m):
    rows = []
    for label, p in PERIODS:
        sub = m[m.index >= p] if isinstance(p, str) else (m.iloc[p:] if p else m)
        for name, col in [("Unlevered (100% invested)", "unlevered"), ("Dial (0-200%)", "dial")]:
            rows.append({"Period": label, "Strategy": name, **_stats(sub[col], sub["rf"]),
                         "Avg exposure": sub["exposure"].mean() if col == "dial" else 1.0,
                         "Months levered (>100%)": int((sub["exposure"] > 1).sum()) if col == "dial" else 0,
                         "Months de-risked (<100%)": int((sub["exposure"] < 1).sum()) if col == "dial" else 0})
    return pd.DataFrame(rows)


NOTES = [
    "Unlevered = the strategy always 100% invested (Step Nine 'Portfolio' return). Dial = exposure 0% to 200%.",
    "Exposure in a month = the dial reading at the last trading day of the previous month; held fixed within the month.",
    "Below 100% the rest earns the T-bill rate (yfinance ^IRX). Above 100% the borrowed part pays T-bill + 0.42% (a fixed assumption).",
    "Monthly returns, gross of trading costs. Drawdowns use month-end values, so intra-month dips (e.g. March 2020) are understated.",
]


def _fmt_table(table):
    """Human-readable strings for the report table (period label shown once per group)."""
    rows, last = [], None
    for _, r in table.iterrows():
        rows.append([r["Period"] if r["Period"] != last else "", r["Strategy"],
                     f"{r['Ann. return (CAGR)']:.1%}", f"{r['Ann. volatility']:.1%}",
                     f"{r['Sharpe (excess over T-bill)']:.2f}", f"{r['Max drawdown']:.1%}",
                     f"{r['Worst month']:.1%}", f"${r['Growth of $1']:,.1f}",
                     f"{r['Avg exposure']:.0%}"])
        last = r["Period"]
    cols = ["Period", "Strategy", "Return / yr", "Volatility / yr", "Sharpe", "Worst drawdown",
            "Worst month", "$1 grows to", "Avg exposure"]
    return cols, rows


def plot_performance(m, out_pdf, title, table):
    """One-page light-mode PDF report: growth of $1 (log), drawdown, and the stats table."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(11, 8.5), facecolor="white")
    gs = fig.add_gridspec(3, 1, height_ratios=[3.0, 1.4, 2.6], hspace=0.38,
                          left=0.07, right=0.97, top=0.90, bottom=0.03)
    a1, a2, a3 = fig.add_subplot(gs[0]), fig.add_subplot(gs[1], sharex=None), fig.add_subplot(gs[2])
    a2.sharex(a1)
    for col, name, color in [("unlevered", "Unlevered (100% invested)", "#7f8c8d"), ("dial", "Dial (0-200%)", "#1f4e79")]:
        cum = (1 + m[col]).cumprod()
        a1.plot(cum.index, cum, color=color, lw=1.4, label=f"{name}: $1 grows to ${cum.iloc[-1]:,.1f}")
        a2.plot(cum.index, cum / cum.cummax() - 1, color=color, lw=1.0)
    a1.set_yscale("log")
    a1.set_ylabel("Growth of $1 (log scale)")
    a1.legend(loc="upper left", fontsize=9)
    a1.grid(alpha=0.3, which="both")
    a2.set_ylabel("Drawdown from peak")
    a2.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
    a2.grid(alpha=0.3)
    fig.suptitle(f"{title}\nUnlevered vs exposure dial, {m.index[0]:%b %Y} to {m.index[-1]:%b %Y}",
                 fontsize=13, fontweight="bold", y=0.975)
    a3.axis("off")
    cols, rows = _fmt_table(table)
    tb = a3.table(cellText=rows, colLabels=cols, loc="upper center", cellLoc="center",
                  colWidths=[0.17, 0.19, 0.08, 0.10, 0.07, 0.10, 0.09, 0.09, 0.09])
    tb.auto_set_font_size(False)
    tb.set_fontsize(8)
    tb.scale(1, 1.25)
    for (r, c), cell in tb.get_celld().items():
        cell.set_edgecolor("#bbbbbb")
        if r == 0:
            cell.set_facecolor("#e8edf3")
            cell.set_text_props(fontweight="bold")
        else:
            is_dial = rows[r - 1][1].startswith("Dial")
            cell.set_facecolor("#f2f7fc" if is_dial else "white")
            if is_dial:
                cell.set_text_props(fontweight="bold")
            if c in (0, 1):
                cell.set_text_props(ha="left")
                cell._loc = "left"
    a3.text(0.0, -0.02, "\n".join("- " + n for n in NOTES), transform=a3.transAxes, fontsize=7.5,
            va="top", color="#444444")
    fig.savefig(out_pdf)
    plt.close(fig)


def write_performance_sheets(xw, table, monthly, stats_sheet, monthly_sheet):
    """Write the two Excel sheets with readable formatting (percent formats, widths, header, notes)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    hdr_fill, grey = PatternFill("solid", fgColor="E8EDF3"), PatternFill("solid", fgColor="F2F7FC")
    bold = Font(bold=True)

    # ---- stats sheet
    t = table.rename(columns={"Ann. return (CAGR)": "Return / yr", "Ann. volatility": "Volatility / yr",
                              "Sharpe (excess over T-bill)": "Sharpe", "Max drawdown": "Worst drawdown",
                              "Growth of $1": "$1 grows to", "Months levered (>100%)": "Months levered",
                              "Months de-risked (<100%)": "Months de-risked"})
    t = t[["Period", "Strategy", "Months", "Return / yr", "Volatility / yr", "Sharpe", "Worst drawdown",
           "Worst month", "Best month", "$1 grows to", "Avg exposure", "Months levered", "Months de-risked"]]
    t.to_excel(xw, sheet_name=stats_sheet, index=False)
    ws = xw.sheets[stats_sheet]
    fmts = {"Return / yr": "0.0%", "Volatility / yr": "0.0%", "Sharpe": "0.00", "Worst drawdown": "0.0%",
            "Worst month": "0.0%", "Best month": "0.0%", "$1 grows to": '"$"#,##0.0', "Avg exposure": "0%"}
    names = list(t.columns)
    for c, name in enumerate(names, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = bold, hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = 30 if name == "Period" else (26 if name == "Strategy" else 15)
    ws.row_dimensions[1].height = 32
    for r in range(2, len(t) + 2):
        is_dial = str(ws.cell(row=r, column=2).value).startswith("Dial")
        for c, name in enumerate(names, start=1):
            cell = ws.cell(row=r, column=c)
            if name in fmts:
                cell.number_format = fmts[name]
            if c > 2:
                cell.alignment = Alignment(horizontal="center")
            if is_dial:
                cell.font, cell.fill = bold, grey
    ws.freeze_panes = "C2"
    for i, note in enumerate(NOTES):
        ws.cell(row=len(t) + 3 + i, column=1, value=note).font = Font(italic=True, color="555555")

    # ---- monthly sheet
    mm = monthly.reset_index(names="Month")[["Month", "unlevered", "exposure", "rf", "dial"]]
    mm.columns = ["Month", "Unlevered return", "Dial exposure held", "T-bill return", "Dial return"]
    mm.to_excel(xw, sheet_name=monthly_sheet, index=False)
    ws = xw.sheets[monthly_sheet]
    for c in range(1, 6):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = bold, hdr_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = 18
    for r in range(2, len(mm) + 2):
        ws.cell(row=r, column=1).number_format = "mmm-yyyy"
        ws.cell(row=r, column=2).number_format = "0.00%"
        ws.cell(row=r, column=3).number_format = "0%"
        ws.cell(row=r, column=4).number_format = "0.00%"
        ws.cell(row=r, column=5).number_format = "0.00%"
    ws.freeze_panes = "B2"
